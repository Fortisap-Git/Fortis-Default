#!/usr/bin/env python3
"""Phase A in one command: download every FYI document in parallel and extract
its text, cached by document id so a re-run costs nothing.

    python scripts/gather.py manifest.json work/<client>/

manifest.json is a list produced from fyi_find_documents + fyi_download_document:
[
  {"id": "<fyi doc uuid>", "name": "2026 ATO Pre-filling report - Jane Smith.pdf",
   "url": "<presigned S3 url from fyi_download_document>"},
  ...
]

For each entry the script writes  <workdir>/docs/<id>__<safe name>  (the file) and
<workdir>/text/<id>.txt  (extracted text) and skips both if they already exist.
PDF text: pdftotext -layout if on PATH, else PyMuPDF, else pypdf. Excel: every
sheet's cells as "A1 = value" lines. Word: python-docx if installed.

Prints a summary table (id, name, pages/chars, extractor) and writes
<workdir>/gather_index.json for the build step.

Presigned URLs live ~15 minutes: run this immediately after fetching them.
"""
import json
import re
import shutil
import subprocess
import sys
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path


def safe(name):
    return re.sub(r"[^A-Za-z0-9._ -]+", "_", name)[:120]


def download(entry, docs):
    dest = docs / f"{entry['id']}__{safe(entry['name'])}"
    if dest.exists() and dest.stat().st_size > 0:
        return dest, "cached"
    if not entry.get("url"):
        return None, "no url"
    req = urllib.request.Request(entry["url"], headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as r, open(dest, "wb") as f:
        shutil.copyfileobj(r, f)
    return dest, "downloaded"


def pdf_text(path):
    if shutil.which("pdftotext"):
        p = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True)
        if p.returncode == 0 and p.stdout.strip():
            return p.stdout.decode("utf-8", "replace"), "pdftotext"
    try:
        import fitz
        doc = fitz.open(path)
        return "\n\f".join(pg.get_text("text") for pg in doc), "pymupdf"
    except Exception:
        pass
    from pypdf import PdfReader
    return "\n\f".join((pg.extract_text() or "") for pg in PdfReader(str(path)).pages), "pypdf"


def xlsx_text(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out.append(f"{c.coordinate} = {c.value}")
    return "\n".join(out), "openpyxl"


def docx_text(path):
    try:
        import docx
    except ImportError:
        return "", "python-docx not installed"
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts), "python-docx"


def extract(path, textdir, doc_id):
    tp = textdir / f"{doc_id}.txt"
    if tp.exists():
        return tp, "cached"
    ext = path.suffix.lower()
    if ext == ".pdf":
        txt, how = pdf_text(path)
    elif ext in (".xlsx", ".xlsm", ".xls"):
        txt, how = xlsx_text(path)
    elif ext == ".docx":
        txt, how = docx_text(path)
    elif ext in (".csv", ".txt"):
        txt, how = path.read_text(errors="replace"), "raw"
    else:
        txt, how = "", f"unsupported {ext}"
    tp.write_text(txt, encoding="utf-8")
    return tp, how


def one(entry, docs, textdir):
    try:
        path, dl = download(entry, docs)
        if not path:
            return {**entry, "status": dl}
        tp, how = extract(path, textdir, entry["id"])
        n = len(tp.read_text(encoding="utf-8", errors="replace"))
        return {**entry, "file": str(path), "text": str(tp), "chars": n,
                "status": f"{dl}/{how}"}
    except Exception as e:  # keep going -- one bad doc must not kill the gather
        return {**entry, "status": f"ERROR {e}"}


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    manifest = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    work = Path(sys.argv[2])
    docs, textdir = work / "docs", work / "text"
    docs.mkdir(parents=True, exist_ok=True)
    textdir.mkdir(parents=True, exist_ok=True)
    with ThreadPoolExecutor(max_workers=8) as ex:
        results = list(ex.map(lambda e: one(e, docs, textdir), manifest))
    (work / "gather_index.json").write_text(json.dumps(results, indent=1), encoding="utf-8")
    w = max(len(r["name"]) for r in results) if results else 10
    for r in results:
        print(f"{r['id'][:8]}  {r['name']:<{w}}  {r.get('chars', 0):>8,}  {r['status']}")
    bad = [r for r in results if r["status"].startswith(("ERROR", "no url"))]
    print(f"\n{len(results)} documents, {len(bad)} problems -> {work / 'gather_index.json'}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
