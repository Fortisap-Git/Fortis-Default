#!/usr/bin/env python3
"""Compare a working ITR workpaper against the bundled master template.

Reports structural and formatting drift only. Values and formulas are ignored --
those are supposed to change; fonts, fills, number formats, widths, sheet order
and sheet visibility are not.

    python scripts/check_template_fidelity.py working.xlsm \
        assets/2026_ITR_Workpaper_-_Individual_Name.xlsm

Exit code 0 = clean, 1 = drift found, 2 = could not run.

Renames of the "1 Client name" / "2 Client name" tabs are expected and are matched
back to their master counterparts. Sheets in the working file with no counterpart
are treated as clones: they are matched to the master sheet they were cloned from
when the name makes that obvious (e.g. "Rental Property - 12 Smith St" ->
"Rental Property"), otherwise reported as unmatched for a human to confirm.
"""

import sys
import re
from collections import Counter

try:
    import openpyxl
except ImportError:  # pragma: no cover
    print("openpyxl is required", file=sys.stderr)
    sys.exit(2)

CLIENT_TAB = re.compile(r"^(Summary|Deductions) - (.+)$")
REVIEW_TABS = ("Review Notes", "Queries")   # content tabs: body-row styles are the preparer's


def norm(name, master_names, taken):
    """Map a renamed client tab back to its master counterpart, or None.
    'Summary - Jane Smith' -> 'Summary - 1 Client name ' (first unmatched),
    the second distinct Summary tab -> '2 Client name '."""
    m = CLIENT_TAB.match(name)
    if not m:
        return None
    for n in ("1", "2"):
        cand = f"{m.group(1)} - {n} Client name "
        if cand in master_names and cand not in taken:
            return cand
    return None


def clone_parent(name, master_names):
    for mn in master_names:
        if name.startswith(mn.rstrip()) and name != mn:
            return mn
    return None


def cell_style(c):
    """Style tuple + a kind flag. Font name/size/bold/italic/colour, fill, number format.
    kind: 'num' (numeric or formula -- every component matters), 'text', 'empty', 'link'.
    Text/empty cells compare without number format (it is inert there); link cells also
    compare without colour (link colour/underline is the one permitted restyle)."""
    f = c.font
    fill = c.fill
    v = c.value
    if bool(c.hyperlink) or (isinstance(v, str) and v.startswith("=HYPERLINK(")):
        kind = "link"
    elif v is None:
        kind = "empty"
    elif isinstance(v, str) and not v.startswith("="):
        kind = "text"
    else:
        kind = "num"
    return (
        f.name,
        f.sz,
        bool(f.b),
        bool(f.i),
        getattr(f.color, "rgb", None) if f.color is not None else None,
        fill.patternType,
        getattr(fill.fgColor, "rgb", None) if fill is not None else None,
        c.number_format,
        kind,
    )


def style_known(s, master_profile):
    """Is working style s acceptable given the master sheet's styles?"""
    core = s[:7]
    for m in master_profile:
        if m[:7] == core and (s[8] != "num" or m[7] == s[7]):
            return True
        if s[8] == "link" and m[:4] == s[:4] and m[5:7] == s[5:7]:
            return True
    return False


def sheet_style_profile(ws, max_row=400, max_col=30):
    prof = Counter()
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_row),
                            max_col=min(ws.max_column, max_col)):
        for c in row:
            prof[cell_style(c)] += 1
    return prof


def widths(ws):
    return {k: round(v.width, 1) for k, v in ws.column_dimensions.items()
            if v.width is not None}


def main():
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    work_path, master_path = sys.argv[1], sys.argv[2]

    wb = openpyxl.load_workbook(work_path, keep_vba=True)
    mb = openpyxl.load_workbook(master_path, keep_vba=True)

    issues, notes = [], []
    master_names = mb.sheetnames

    # --- sheet inventory -------------------------------------------------
    mapping = {}          # working sheet -> master sheet
    clones = {}
    for name in wb.sheetnames:
        if name in master_names:
            mapping[name] = name
        elif norm(name, master_names, set(mapping.values())):
            mapping[name] = norm(name, master_names, set(mapping.values()))
        else:
            parent = clone_parent(name, master_names)
            if parent:
                clones[name] = parent
                mapping[name] = parent
            else:
                issues.append(f"SHEET: '{name}' has no counterpart in the master "
                              f"(new tab -- confirm it is intended)")

    covered = set(mapping.values())
    for name in master_names:
        if name not in covered:
            issues.append(f"SHEET: master sheet '{name}' is missing from the working file")

    # --- sheet order (of the non-clone sheets) ---------------------------
    seq = [mapping[n] for n in wb.sheetnames if n in mapping and n not in clones]
    seq_master = [n for n in master_names if n in set(seq)]
    if seq != seq_master:
        issues.append("ORDER: sheet order differs from the master\n"
                      f"        master:  {seq_master}\n"
                      f"        working: {seq}")

    # --- visibility, widths, merges, styles ------------------------------
    for wname, mname in mapping.items():
        ws, ms = wb[wname], mb[mname]
        tag = wname if wname == mname else f"{wname} (<- {mname})"

        if ws.sheet_state != ms.sheet_state:
            if ms.sheet_state == "hidden" and ws.sheet_state == "visible":
                notes.append(f"unhidden: '{tag}' (allowed when the return uses it)")
            else:
                issues.append(f"VISIBILITY: '{tag}' is {ws.sheet_state}, "
                              f"master is {ms.sheet_state}")

        ww, mw = widths(ws), widths(ms)
        changed = {k: (mw[k], ww[k]) for k in mw if k in ww and mw[k] != ww[k]}
        if changed:
            issues.append(f"WIDTH: '{tag}' column widths changed "
                          f"{dict(list(changed.items())[:8])}")

        wmerge = {str(r) for r in ws.merged_cells.ranges}
        mmerge = {str(r) for r in ms.merged_cells.ranges}
        if wmerge != mmerge:
            lost, added = sorted(mmerge - wmerge)[:8], sorted(wmerge - mmerge)[:8]
            issues.append(f"MERGE: '{tag}' merged ranges differ (lost {lost}, added {added})")

        if mname in REVIEW_TABS:
            continue          # the review record / query list are written content, not layout
        wp, mp = sheet_style_profile(ws), sheet_style_profile(ms)
        new_styles = [s for s in wp if not style_known(s, mp)]
        for s in new_styles[:6]:
            font, size, bold, ital, colour, pattern, fill, numfmt, kind = s
            issues.append(
                f"STYLE: '{tag}' uses a style not in the master -- "
                f"font={font} {size} bold={bold} colour={colour} "
                f"fill={fill} numfmt={numfmt!r} ({wp[s]} cells)")
        if len(new_styles) > 6:
            issues.append(f"STYLE: '{tag}' has {len(new_styles) - 6} further new styles")

    # --- macros ----------------------------------------------------------
    import zipfile
    if "xl/vbaProject.bin" not in zipfile.ZipFile(work_path).namelist():
        issues.append("MACROS: xl/vbaProject.bin is missing -- reload with keep_vba=True")

    if clones:
        print(f"Cloned tabs matched to their master parent: {clones}\n")
    for n in notes:
        print(" i", n)
    if issues:
        print(f"{len(issues)} fidelity issue(s):\n")
        for i in issues:
            print(" -", i)
        return 1
    print("Clean: structure, fonts, colours, number formats and widths match the master.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
