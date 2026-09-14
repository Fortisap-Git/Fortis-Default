"""Low-level helpers for the Fortis ITR workpaper (openpyxl).

The build engine (build_workpaper.py) is the only thing that should normally
call these. They encode the traps hit on real client runs:

  * keep_vba=True or the .xlsm loses its macros
  * sheet renames update nothing -- formula strings and hyperlink locations
    must be patched by hand
  * the master ships 288 defined names pointing at dead external workbooks;
    delete them and clear the external-link list or recalculation refuses
  * never replace a cell's Font -- copy it and change only colour/underline,
    or the fidelity check fails (Arial 10 silently becomes Calibri 11)
  * totals are formulas or in-cell sums, never Python-computed constants
"""
from copy import copy
import re

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.styles.colors import Color
from openpyxl.worksheet.hyperlink import Hyperlink

LINK_RGB = "FF0070C0"          # the master's own link blue -- introduce no other colour
FYI_URL = "https://go.fyi.app/search/0/{entity_id}/0/documents/{doc_uuid}/preview"
EXT_RE = re.compile(r"\[\d+\]")


def load_template(path):
    """Always keep_vba=True -- .xlsm loses its macros otherwise."""
    return openpyxl.load_workbook(path, keep_vba=True)


def patch_renames(wb, renames):
    """After ws.title changes, openpyxl updates NOTHING. Patch formula strings,
    hyperlink locations and defined names across the whole workbook.
    renames: {old_sheet_name: new_sheet_name}
    """
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("="):
                    nv = v
                    for old, new in renames.items():
                        nv = nv.replace(f"'{old}'", f"'{new}'")
                    if nv != v:
                        c.value = nv
                if c.hyperlink and c.hyperlink.location:
                    loc = c.hyperlink.location
                    nl = loc
                    for old, new in renames.items():
                        nl = nl.replace(f"'{old}'", f"'{new}'")
                    if nl != loc:
                        c.hyperlink.location = nl
    for name, dn in list(wb.defined_names.items()):
        t = dn.attr_text or ""
        nt = t
        for old, new in renames.items():
            nt = nt.replace(f"'{old}'", f"'{new}'")
        if nt != t:
            dn.attr_text = nt


def neutralize_external_links(wb, external_names=None):
    """Delete defined names that point at external workbooks and clear the
    external-link list. No cell formula in the master depends on them (the
    template metadata proves it), so nothing needs baking in.
    external_names: list from template_meta.json; recomputed if None.
    Returns the number of names removed.
    """
    removed = 0
    if external_names is None:
        external_names = [k for k, v in wb.defined_names.items()
                          if v.attr_text and EXT_RE.search(v.attr_text)]
        for ws in wb.worksheets:
            external_names += [f"{ws.title}!{k}" for k, v in ws.defined_names.items()
                               if v.attr_text and EXT_RE.search(v.attr_text)]
    for name in external_names:
        if "!" in name:
            sheet, local = name.split("!", 1)
            if sheet in wb.sheetnames and local in wb[sheet].defined_names:
                del wb[sheet].defined_names[local]
                removed += 1
        elif name in wb.defined_names:
            del wb.defined_names[name]
            removed += 1
    try:
        wb._external_links = []
    except Exception:
        pass
    return removed


def clone_sheet(wb_target, ws_src, new_name=None, after=None):
    """Clone a sheet (from this or another workbook): values, formulas, styles,
    merges, column widths, row heights. Hyperlink objects and embedded images
    DO NOT copy -- the engine recreates links from the spec.
    """
    idx = None
    if after and after in wb_target.sheetnames:
        idx = wb_target.sheetnames.index(after) + 1
    tgt = wb_target.create_sheet(new_name or ws_src.title, index=idx)
    for row in ws_src.iter_rows():
        for c in row:
            if type(c).__name__ == "MergedCell":
                continue
            nc = tgt.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                nc.font = copy(c.font)
                nc.fill = copy(c.fill)
                nc.border = copy(c.border)
                nc.alignment = copy(c.alignment)
                nc.number_format = c.number_format
                nc.protection = copy(c.protection)
    for rng in ws_src.merged_cells.ranges:
        tgt.merge_cells(str(rng))
    for col, dim in ws_src.column_dimensions.items():
        tgt.column_dimensions[col].width = dim.width
        tgt.column_dimensions[col].hidden = dim.hidden
    for r, dim in ws_src.row_dimensions.items():
        tgt.row_dimensions[r].height = dim.height
    tgt.sheet_state = "visible"
    return tgt


def roll_comparatives(ws, ws_prior_values, first_row, last_row,
                      cur_cols=("E", "F", "G"), cmp_cols=("B", "C", "D")):
    """Prior-year current columns become this year's comparatives, then the
    current columns are CLEARED (constants only -- formulas stay).

    last_row must cover the ENTIRE used range including total/ownership/control
    rows -- a stale prior-year constant left in an ownership row was the worst
    bug of the first run.
    """
    for r in range(first_row, last_row + 1):
        for tc, sc in zip(cmp_cols, cur_cols):
            cell = ws[f"{tc}{r}"]
            if isinstance(cell.value, str) and cell.value.startswith("="):
                continue
            cell.value = ws_prior_values[f"{sc}{r}"].value
        for cc in cur_cols:
            cell = ws[f"{cc}{r}"]
            if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None


def _font_with(cell, **over):
    """The cell's own font with a few attributes overridden -- never a bare
    Font(), which would reset Arial 10 to Calibri 11 and fail fidelity."""
    f = cell.font
    kw = dict(name=f.name, sz=f.sz, b=f.b, i=f.i, vertAlign=f.vertAlign,
              strike=f.strike, color=f.color, underline=f.underline,
              family=f.family, charset=f.charset, scheme=f.scheme)
    kw.update(over)
    return Font(**kw)


def link_font(cell):
    return _font_with(cell, color=Color(rgb=LINK_RGB), underline="single")


def link(ws, coord, location, text=None, restyle=True):
    """Internal cross-reference hyperlink, firm convention: text 'Sheet - Cell',
    location 'Sheet'!$Cell$."""
    c = ws[coord]
    if text is not None:
        c.value = text
    c.hyperlink = Hyperlink(ref=coord, location=location)
    if restyle and not c.font.u:      # already-styled link columns stay as they are
        c.font = link_font(c)


def fyi_link(ws, coord, entity_id, doc_uuid, label, restyle=True):
    """FYI stable document link: =HYPERLINK() to the go.fyi.app preview."""
    url = FYI_URL.format(entity_id=entity_id, doc_uuid=doc_uuid)
    c = ws[coord]
    c.value = f'=HYPERLINK("{url}", "{label}")'
    if restyle and not c.font.u:
        c.font = link_font(c)


def bold(cell):
    cell.font = _font_with(cell, b=True)


def wrap(cell):
    a = cell.alignment
    cell.alignment = Alignment(horizontal=a.horizontal, vertical="top", wrap_text=True)


def copy_row_style(ws, src_row, dst_row, max_col=None):
    max_col = max_col or ws.max_column
    for col in range(1, max_col + 1):
        s, d = ws.cell(src_row, col), ws.cell(dst_row, col)
        if s.has_style:
            d.font, d.fill, d.border = copy(s.font), copy(s.fill), copy(s.border)
            d.alignment, d.number_format = copy(s.alignment), s.number_format
    if ws.row_dimensions[src_row].height:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def validate_links(wb):
    """Every internal hyperlink must target an existing sheet. Returns broken refs."""
    names = set(wb.sheetnames)
    broken = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.hyperlink and c.hyperlink.location:
                    sheet = c.hyperlink.location.split("!")[0].strip("'")
                    if sheet not in names:
                        broken.append(f"{ws.title}!{c.coordinate} -> {c.hyperlink.location}")
    return broken


def verify_vba(path):
    import zipfile
    return "xl/vbaProject.bin" in zipfile.ZipFile(path).namelist()


# --------------------------------------------------------------------------
# Phase E -- the review record lives on the workpaper's Review Notes tab.
# Organise the register the way a reviewer works the file: by NATURE, in ITR
# label order, not by severity and not by which lens found it.
# --------------------------------------------------------------------------

SECTIONS = (
    "INCOME — by label (10 interest, 11 dividends, 13 distributions, 18 CGT, 20 foreign)",
    "RENTAL PROPERTY — in rental schedule order",
    "DEDUCTIONS — D1 to D15 in label order",
    "OTHER — offsets, levies, disclosures",
    "WORKPAPER MECHANICS — fixed in this file",
    "OUTSIDE THIS FILE",
)

REGISTER_COLS = ("#", "Sheet", "Cell", "Review point", "Action required",
                 "Status", "Tax effect", "Source doc")


def write_review_notes(ws, summary, sections, start_row=8):
    """Write the review record: manager summary block, then the register grouped
    by nature. Uses the tab's own fonts (bold for headings) -- no fills, no
    column widths, nothing the fidelity check would flag.

    summary:  {"bottom_line","waiting","fixed","outside"} -- plain English.
    sections: ordered [(section_title, [finding, ...]), ...]
    finding:  {sheet, cell, point, action, status, impact, source}
    Returns {(section_index, finding_index): row} for flag_cell.
    """
    for row in ws.iter_rows(min_row=start_row, max_row=max(ws.max_row, start_row)):
        for c in row:
            c.value, c.comment = None, None

    r = start_row
    ws.cell(row=r, column=1, value="REVIEW SUMMARY"); bold(ws.cell(row=r, column=1))
    r += 1
    for label, key in (("Bottom line", "bottom_line"),
                       ("Waiting on the client", "waiting"),
                       ("Fixed during review", "fixed"),
                       ("Outside this file", "outside")):
        ws.cell(row=r, column=1, value=label); bold(ws.cell(row=r, column=1))
        ws.cell(row=r, column=2, value=summary.get(key, "")); wrap(ws.cell(row=r, column=2))
        r += 1
    r += 1

    rows, n = {}, 0
    for si, (title, findings) in enumerate(sections):
        if not findings:
            continue
        ws.cell(row=r, column=1, value=title); bold(ws.cell(row=r, column=1))
        r += 1
        for col, h in enumerate(REGISTER_COLS, start=1):
            ws.cell(row=r, column=col, value=h); bold(ws.cell(row=r, column=col))
        r += 1
        for fi, f in enumerate(findings):
            n += 1
            ws.cell(row=r, column=1, value=n)
            ws.cell(row=r, column=2, value=f.get("sheet", ""))
            ws.cell(row=r, column=3, value=f.get("cell", ""))
            ws.cell(row=r, column=4, value=f.get("point", "")); wrap(ws.cell(row=r, column=4))
            ws.cell(row=r, column=5, value=f.get("action", "")); wrap(ws.cell(row=r, column=5))
            ws.cell(row=r, column=6, value=f.get("status", ""))
            ws.cell(row=r, column=7, value=f.get("impact", ""))
            ws.cell(row=r, column=8, value=f.get("source", "")); wrap(ws.cell(row=r, column=8))
            rows[(si, fi)] = r
            r += 1
        r += 1
    return rows


def flag_cell(wb, sheet, coord, register_row):
    """Cell comment at a flagged cell pointing back at its Review Notes row."""
    from openpyxl.comments import Comment
    if sheet not in wb.sheetnames or not coord:
        return
    wb[sheet][coord.split(",")[0].strip()].comment = Comment(
        f"Review Notes row {register_row}", "Review")
