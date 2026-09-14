#!/usr/bin/env python3
"""Build the ITR workpaper from a JSON fill spec. Deterministic, idempotent,
runs in a few seconds. The model writes the SPEC (data); this script does the
Excel work. Never author a per-client Python build script again.

    python scripts/build_workpaper.py spec.json            # build
    python scripts/build_workpaper.py spec.json --dry-run  # validate spec only

Spec format: see references/fill-spec.md. Minimal example:

{
  "client":   {"name": "Jane Smith", "entity_id": 12345, "year": 2026},
  "template": null,                       # null -> bundled master
  "prior_workpaper": "work/2025 ITR Workpaper - Jane Smith.xlsm",   # or null
  "output":   "work/2026 ITR Workpaper - Jane Smith.xlsm",   # must be .xlsm (macro-enabled)
  "unhide":   ["Share Register"],
  "clones":   [{"prior_sheet": "Rental Property - 12 Smith St", "as": "Rental Property - 12 Smith St",
                "after": "Rental Property", "roll": {"first_row": 7, "last_row": 56}}],
  "sources":  [{"id": "prefill", "doc": "<fyi-uuid>", "label": "2026 ATO Pre-filling report"}],
                                          # cited per figure; links land in the Hyperlink column only
  "cells":    [{"sheet": "Summary", "cell": "H16", "value": 1234.56, "source": "prefill",
                "remark": "Per pre-fill; re-run before lodgement"}],
  "queries":  [{"issue": "Interest", "description": "...", "reference": {"sheet": "Summary", "cell": "H16"},
                "client_reply": null, "reply": null}],   # replies filled in on an update run
  "run":      {"round": 1, "built": "2026-09-14"},       # bump each update run; the spec is the state
  "review_notes": ["Ownership split 50/50 per prior return schedule"],
  "rev":      [],                          # Phase E overrides, same shape as cells
  "review":   null                         # Phase E {summary, sections} -> Review Notes register
}

Sheet aliases accepted everywhere: "Summary" / "Deductions" (client 1 tabs after rename),
"Summary2" / "Deductions2" (client 2), otherwise the exact sheet name.
"""
import json
import re
import shutil
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import wp_helpers as h  # noqa: E402

HERE = Path(__file__).resolve().parent
ASSETS = HERE.parent / "assets"
MASTER = ASSETS / "2026_ITR_Workpaper_-_Individual_Name.xlsm"
META = json.loads((ASSETS / "template_meta.json").read_text(encoding="utf-8"))

S1, D1 = "Summary - 1 Client name ", "Deductions - 1 Client name "
S2, D2 = "Summary - 2 Client name ", "Deductions - 2 Client name "


class Build:
    def __init__(self, spec):
        self.spec = spec
        self.warnings, self.log = [], []
        self.alias = {}
        self.entity_id = spec["client"].get("entity_id")
        self.sources = {s["id"]: s for s in spec.get("sources", [])}
        self.query_row_of = {}

    # ---------------------------------------------------------------- utils
    def warn(self, msg):
        self.warnings.append(msg)

    def sheet(self, name):
        real = self.alias.get(name, name)
        if real not in self.wb.sheetnames:
            raise KeyError(f"sheet '{name}' not in workbook (aliases: {sorted(self.alias)})")
        return self.wb[real]

    def landmarks(self, ws):
        t = ws.title
        L = META["landmarks"]
        if t.startswith("Summary - "):
            return L["summary"]
        if t.startswith("Deductions - "):
            return L["deductions"]
        if t.startswith("Rental Property"):
            return L["rental"]
        return {}

    @staticmethod
    def is_formula(v):
        return isinstance(v, str) and v.startswith("=")

    # ---------------------------------------------------------------- steps
    def load(self):
        spec = self.spec
        tpl = spec.get("template") or MASTER
        tpl = Path(tpl)
        if tpl.resolve() == MASTER.resolve():
            self.log.append("template: bundled master")
        else:
            self.log.append(f"template: {tpl}")
        self.wb = h.load_template(tpl)
        self.prior = self.prior_v = None
        if spec.get("prior_workpaper"):
            pp = Path(spec["prior_workpaper"])
            self.prior = openpyxl.load_workbook(pp, keep_vba=True)
            self.prior_v = openpyxl.load_workbook(pp, data_only=True)
            self.log.append(f"prior workpaper: {pp.name}")

    def rename(self):
        c1 = self.spec["client"]["name"]
        c2 = (self.spec.get("client2") or {}).get("name")
        renames = {}
        for old, new in ((S1, f"Summary - {c1}"), (D1, f"Deductions - {c1}")):
            if old in self.wb.sheetnames:
                self.wb[old].title = new
                renames[old] = new
        self.alias.update({"Summary": renames.get(S1, S1), "Deductions": renames.get(D1, D1)})
        if c2:
            for old, new in ((S2, f"Summary - {c2}"), (D2, f"Deductions - {c2}")):
                if old in self.wb.sheetnames:
                    self.wb[old].title = new
                    renames[old] = new
        self.alias.update({"Summary2": renames.get(S2, S2), "Deductions2": renames.get(D2, D2)})
        if renames:
            h.patch_renames(self.wb, renames)
        # client name cells on the client tabs
        ws = self.sheet("Summary")
        ws[META["landmarks"]["summary"]["client_cell"]] = c1
        if c2 and "Summary2" in self.alias and self.alias["Summary2"] in self.wb.sheetnames:
            self.sheet("Summary2")[META["landmarks"]["summary"]["client_cell"]] = c2
        year = self.spec["client"].get("year")
        if year:
            import datetime
            ws["C4"] = datetime.datetime(int(year), 6, 30)
        self.log.append(f"renamed: {list(renames.values())}")

    def neutralize(self):
        n = h.neutralize_external_links(self.wb, META.get("external_names"))
        self.log.append(f"external names removed: {n}")

    def unhide(self):
        for name in self.spec.get("unhide", []):
            self.sheet(name).sheet_state = "visible"
            self.log.append(f"unhidden: {name}")

    def clones(self):
        for cl in self.spec.get("clones", []):
            new = cl["as"]
            if new in self.wb.sheetnames:
                self.warn(f"clone target '{new}' already exists; skipped")
                continue
            if cl.get("prior_sheet"):
                if not self.prior:
                    raise ValueError(f"clone '{new}' needs prior_workpaper")
                src = self.prior[cl["prior_sheet"]]
                tgt = h.clone_sheet(self.wb, src, new, after=cl.get("after"))
                roll = cl.get("roll")
                if roll:
                    h.roll_comparatives(tgt, self.prior_v[cl["prior_sheet"]],
                                        roll["first_row"], roll["last_row"],
                                        tuple(roll.get("cur_cols", ("E", "F", "G"))),
                                        tuple(roll.get("cmp_cols", ("B", "C", "D"))))
                self.log.append(f"cloned from prior: {cl['prior_sheet']} -> {new}"
                                + (" (rolled)" if roll else ""))
            elif cl.get("template_sheet"):
                src = self.sheet(cl["template_sheet"])
                h.clone_sheet(self.wb, src, new, after=cl.get("after") or src.title)
                self.log.append(f"cloned from template: {src.title} -> {new}")
            else:
                raise ValueError(f"clone '{new}' needs prior_sheet or template_sheet")
            # year headers on rental clones
            if new.startswith("Rental Property") and self.spec["client"].get("year"):
                y = int(self.spec["client"]["year"])
                self.wb[new]["B4"], self.wb[new]["E4"] = f"FY {y-1}", f"FY {y}"

    def write_cell(self, e):
        ws = self.sheet(e["sheet"])
        coord = e["cell"]
        cell = ws[coord]
        lm = self.landmarks(ws)
        row = cell.row

        if "formula" in e and e["formula"] is not None:
            f = e["formula"]
            cell.value = f if f.startswith("=") else "=" + f
        elif "value" in e:
            if self.is_formula(cell.value) and not e.get("overwrite_formula"):
                self.warn(f"{ws.title}!{coord}: template formula {cell.value!r} replaced by "
                          f"constant {e['value']!r} (set overwrite_formula:true to silence)")
            cell.value = e["value"]
        elif "link" in e:
            lk = e["link"]
            loc = lk.get("location") or f"'{self.alias.get(lk['sheet'], lk['sheet'])}'!${re.sub(r'(\d+)', r'$\1', lk['cell'])}"
            h.link(ws, coord, loc, lk.get("text"))
        elif "fyi" in e:
            src = self.sources[e["fyi"]] if isinstance(e["fyi"], str) else e["fyi"]
            h.fyi_link(ws, coord, self.entity_id, src["doc"], src["label"])
        elif "clear" in e:
            cell.value = None
        else:
            raise ValueError(f"cell entry needs value/formula/link/fyi/clear: {e}")

        if e.get("number_format"):
            cell.number_format = e["number_format"]

        # source hyperlink in the side column
        src = e.get("source")
        if src:
            sc = e.get("source_cell") or (f"{lm['source_col']}{row}" if lm.get("source_col") else None)
            if not sc:
                self.warn(f"{ws.title}!{coord}: has source but sheet has no source column; "
                          f"give source_cell")
            else:
                if isinstance(src, str):
                    src = self.sources.get(src) or {"text": src}
                if src.get("doc"):
                    h.fyi_link(ws, sc, self.entity_id, src["doc"], src.get("label", "source"))
                elif src.get("sheet"):
                    loc = f"'{self.alias.get(src['sheet'], src['sheet'])}'!${re.sub(r'(\d+)', r'$\1', src['cell'])}"
                    h.link(ws, sc, loc, src.get("text") or f"{src['sheet']} - {src['cell']}")
                else:
                    ws[sc] = src.get("text", "")
        rem = e.get("remark")
        if rem:
            rc = e.get("remark_cell") or (f"{lm['remark_col']}{row}" if lm.get("remark_col") else None)
            if rc:
                ws[rc] = rem
            else:
                self.warn(f"{ws.title}!{coord}: remark given but sheet has no remark column")

    def cells(self, key="cells"):
        n = 0
        for e in self.spec.get(key, []) or []:
            self.write_cell(e)
            n += 1
        self.log.append(f"{key}: {n} written")

    # The Summary tab's "Source Documents" block (J1 heading, J2..J9 and the
    # columns beside it) is deliberately left empty. Firm rule: a source
    # document is linked where the figure it supports sits -- the Hyperlink
    # column (Summary I, Rental H) via a cell's "source" -- never as a bulk
    # index of every document on the file.

    def queries(self):
        qs = self.spec.get("queries", []) or []
        if not qs:
            return
        ws = self.sheet("Queries")
        cols = META["landmarks"]["queries"]["cols"]
        rows = list(META["labels"]["query_rows"])
        last = rows[-1]
        while len(rows) < len(qs):
            last += 1
            h.copy_row_style(ws, rows[-1], last, 6)
            ws[f"{cols['sr']}{last}"] = len(rows) + 1
            rows.append(last)
        for i, q in enumerate(qs):
            r = rows[i]
            ws[f"{cols['issue']}{r}"] = q.get("issue", "")
            ws[f"{cols['description']}{r}"] = q.get("description", "")
            # update runs: what the client came back with, and our response to it
            if q.get("client_reply") is not None:
                ws[f"{cols['client_reply']}{r}"] = q["client_reply"]
            if q.get("reply") is not None:
                ws[f"{cols['reply']}{r}"] = q["reply"]
            ref = q.get("reference")
            if isinstance(ref, dict):
                real = self.alias.get(ref["sheet"], ref["sheet"])
                loc = f"'{real}'!${re.sub(r'(\d+)', r'$\1', ref['cell'])}"
                h.link(ws, f"{cols['reference']}{r}", loc, f"{real.strip()} - {ref['cell']}")
                if q.get("backlink", True):
                    tws = self.sheet(ref["sheet"])
                    lm = self.landmarks(tws)
                    bc = q.get("backlink_cell") or (
                        f"{lm['source_col']}{tws[ref['cell']].row}" if lm.get("source_col") else None)
                    if bc and tws[bc].value in (None, ""):
                        h.link(tws, bc, f"'Queries'!$B${r}", f"Query {i+1}")
            elif ref:
                ws[f"{cols['reference']}{r}"] = str(ref)
            self.query_row_of[i + 1] = r
        # clear any leftover numbered rows below the last query
        for r in rows[len(qs):]:
            for c in ("issue", "description", "reference", "client_reply", "reply"):
                ws[f"{cols[c]}{r}"] = None
        self.log.append(f"queries: {len(qs)}")

    def review_notes(self):
        """Plain judgment rows (Phase B). Skipped when a Phase E register is present,
        because the register must absorb them (one list, no duplicates)."""
        notes = self.spec.get("review_notes", []) or []
        if self.spec.get("review") or not notes:
            return
        ws = self.sheet("Review Notes")
        lm = META["landmarks"]["review_notes"]
        rows = list(META["labels"]["review_rows"])
        last = rows[-1]
        while len(rows) < len(notes):
            last += 1
            h.copy_row_style(ws, rows[-1], last, 4)
            ws[f"{lm['num_col']}{last}"] = len(rows) + 1
            rows.append(last)
        for i, t in enumerate(notes):
            ws[f"{lm['text_col']}{rows[i]}"] = t if isinstance(t, str) else t.get("text", "")
            h.wrap(ws[f"{lm['text_col']}{rows[i]}"])
        self.log.append(f"review notes: {len(notes)}")

    def review(self):
        rv = self.spec.get("review")
        if not rv:
            return
        ws = self.sheet("Review Notes")
        sections = [(s["title"], s.get("findings", [])) for s in rv.get("sections", [])]
        rows = h.write_review_notes(ws, rv.get("summary", {}), sections,
                                    META["landmarks"]["review_notes"]["first_row"])
        flagged = 0
        for (si, fi), r in rows.items():
            f = sections[si][1][fi]
            if f.get("sheet") and f.get("cell"):
                real = self.alias.get(f["sheet"], f["sheet"])
                if real in self.wb.sheetnames:
                    h.flag_cell(self.wb, real, f["cell"], r)
                    flagged += 1
                else:
                    self.warn(f"review finding points at unknown sheet '{f['sheet']}'")
        self.log.append(f"review register: {sum(len(s[1]) for s in sections)} findings, "
                        f"{flagged} cells flagged")

    def save(self):
        out = Path(self.spec["output"])
        if out.suffix.lower() != ".xlsm":
            raise ValueError(f"output must be a macro-enabled .xlsm workbook, got '{out.name}'. "
                             f"The deliverable filed to FYI is this file, macros and all.")
        out.parent.mkdir(parents=True, exist_ok=True)
        if out.resolve() == MASTER.resolve():
            raise ValueError("refusing to overwrite the bundled master")
        self.wb.calculation.fullCalcOnLoad = True   # Excel recalculates on open
        self.wb.save(out)
        if not h.verify_vba(out):
            raise RuntimeError("vbaProject.bin missing after save")
        broken = h.validate_links(self.wb)
        for b in broken:
            self.warn(f"broken internal link: {b}")
        self.log.append(f"saved: {out} ({out.stat().st_size:,} bytes, VBA ok)")
        return out

    def run(self):
        meta = self.spec.get("run") or {}
        if meta:
            self.log.append("run: " + ", ".join(f"{k}={v}" for k, v in meta.items()))
        self.load()
        self.rename()
        self.neutralize()
        self.unhide()
        self.clones()
        self.cells("cells")
        self.queries()
        self.review_notes()
        self.cells("rev")
        self.review()
        return self.save()


def validate(spec):
    errs = []
    for k in ("client", "output"):
        if k not in spec:
            errs.append(f"missing '{k}'")
    if "client" in spec and "name" not in spec["client"]:
        errs.append("client.name missing")
    out = spec.get("output")
    if out and not str(out).lower().endswith(".xlsm"):
        errs.append(f"output '{out}' must end in .xlsm — the workpaper carries VBA and is "
                    f"delivered and filed macro-enabled")
    ids = {s["id"] for s in spec.get("sources", [])}
    for e in spec.get("cells", []) + spec.get("rev", []):
        if "sheet" not in e or "cell" not in e:
            errs.append(f"cell entry missing sheet/cell: {e}")
        if isinstance(e.get("source"), str) and e["source"] not in ids and "doc" not in e:
            errs.append(f"{e.get('sheet')}!{e.get('cell')}: source id '{e['source']}' not in sources")
        if isinstance(e.get("fyi"), str) and e["fyi"] not in ids:
            errs.append(f"{e.get('sheet')}!{e.get('cell')}: fyi id '{e['fyi']}' not in sources")
    return errs


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    spec = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    errs = validate(spec)
    if errs:
        print("SPEC ERRORS:")
        for e in errs:
            print(" -", e)
        return 1
    if "--dry-run" in sys.argv:
        print("spec valid")
        return 0
    b = Build(spec)
    out = b.run()
    for line in b.log:
        print("  ", line)
    if b.warnings:
        print(f"{len(b.warnings)} warning(s):")
        for w in b.warnings:
            print(" !", w)
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
