#!/usr/bin/env python3
"""Generate assets/template_meta.json from the bundled master template.

Run ONCE when the master changes (not per client). The build engine reads the
JSON instead of exploring the workbook, so a client build never has to
rediscover where label 10 lives or which defined names point at dead
external workbooks.

    python scripts/build_template_meta.py [assets/2026_ITR_Workpaper_-_Individual_Name.xlsm]

Contents:
  sheets           - name, state, order (fidelity baseline)
  external_names   - defined names whose target is an external workbook (deleted on build)
  baked_values     - {sheet: {cell: cached_value}} for any formula that depended on an
                     external name (baked in so recalculation never sees #NAME?)
  landmarks        - per-sheet column roles (value / source / remark columns) and
                     header rows the engine relies on
  labels           - Summary rows by ITR label + description, Deductions block rows,
                     Rental Property rows by line description, Queries / Review Notes
                     numbered rows
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
DEFAULT = HERE.parent / "assets" / "2026_ITR_Workpaper_-_Individual_Name.xlsm"
OUT = HERE.parent / "assets" / "template_meta.json"

SUMMARY = "Summary - 1 Client name "
DEDUCTIONS = "Deductions - 1 Client name "


def s(v):
    return v.strip() if isinstance(v, str) else v


def main(path=DEFAULT):
    wb = openpyxl.load_workbook(path, keep_vba=True)
    wbv = openpyxl.load_workbook(path, data_only=True)

    meta = {"source": Path(path).name, "sheets": [], "external_names": [],
            "baked_values": {}, "landmarks": {}, "labels": {}}

    for ws in wb.worksheets:
        meta["sheets"].append({"name": ws.title, "state": ws.sheet_state})

    # --- external defined names ------------------------------------------
    ext_names = [k for k, v in wb.defined_names.items()
                 if v.attr_text and re.search(r"\[\d+\]", v.attr_text)]
    meta["external_names"] = sorted(ext_names)
    for ws in wb.worksheets:
        for k, v in list(ws.defined_names.items()):
            if v.attr_text and re.search(r"\[\d+\]", v.attr_text):
                meta["external_names"].append(f"{ws.title}!{k}")
    meta["external_link_count"] = len(wb._external_links)

    # formulas that reference an external name or an external workbook directly
    pat = None
    if ext_names:
        pat = re.compile(r"(?<![A-Za-z0-9_])(" + "|".join(map(re.escape, ext_names)) + r")(?![A-Za-z0-9_])")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                # strip quoted sheet names / string literals before name matching,
                # otherwise a name like "name" matches inside 'Summary - 1 Client name '
                bare = re.sub(r"'[^']*'|\"[^\"]*\"", "", v)
                if re.search(r"\[\d+\]", v) or (pat and pat.search(bare)):
                    meta["baked_values"].setdefault(ws.title, {})[c.coordinate] = \
                        wbv[ws.title][c.coordinate].value

    # --- landmarks ---------------------------------------------------------
    meta["landmarks"] = {
        "summary": {"header_row": 10, "label_col": "B", "desc_col": "C",
                    "credit_col": "G", "value_col": "H", "source_col": "I",
                    "remark_col": "J", "signoff_col": "K", "py_marker": "I8",
                    "source_index_anchor": "J1", "client_cell": "C3"},
        "deductions": {"header_row": 9, "desc_col": "B", "items_col": "C",
                       "total_col": "D", "client_cell": "B3"},
        "rental": {"py_cols": ["B", "C", "D"], "cy_cols": ["E", "F", "G"],
                   "total_col": "E", "agency_col": "F", "owner_col": "G",
                   "source_col": "H", "remark_col": "I", "year_row": 4,
                   "first_data_row": 7, "control_row": 56},
        "queries": {"header_row": 7, "cols": {"sr": "A", "issue": "B", "description": "C",
                    "reference": "D", "client_reply": "E", "reply": "F"}},
        "review_notes": {"first_row": 8, "num_col": "B", "text_col": "C"},
        "share_register": {"header_rows": [11, 12], "first_data_row": 14},
        "foreign_income": {"header_row": 9, "first_data_row": 10},
        "dividends": {"header_row": 7, "first_data_row": 13},
        "home_office": {"client1_cols": ["A", "B", "C"], "client2_cols": ["E", "F", "G"]},
    }

    # --- Summary labels ----------------------------------------------------
    ws = wb[SUMMARY]
    rows = []
    for r in range(11, ws.max_row + 1):
        lab, desc = s(ws.cell(r, 2).value), s(ws.cell(r, 3).value)
        if lab is None and desc is None:
            continue
        rows.append({"row": r, "label": str(lab) if lab is not None else None, "desc": desc,
                     "formula": ws.cell(r, 8).value if isinstance(ws.cell(r, 8).value, str)
                     and str(ws.cell(r, 8).value).startswith("=") else None})
    meta["labels"]["summary"] = rows

    # --- Deductions blocks -------------------------------------------------
    ws = wb[DEDUCTIONS]
    blocks, cur = [], None
    for r in range(10, ws.max_row + 1):
        a, b = s(ws.cell(r, 1).value), s(ws.cell(r, 2).value)
        if a and re.match(r"^D\d+", str(a)):
            cur = {"label": str(a).split(" -")[0], "title": a, "row": r, "lines": [], "total_row": None}
            blocks.append(cur)
        if cur is None:
            continue
        if b:
            if str(b).lower().startswith("total"):
                cur["total_row"] = r
            elif r != cur["row"] or a is None:
                cur["lines"].append({"row": r, "desc": b})
            else:
                cur["lines"].append({"row": r, "desc": b})
        if b and str(b).lower().startswith("total deductions"):
            meta["labels"]["deductions_total_row"] = r
    meta["labels"]["deductions"] = blocks

    # --- Rental Property lines ----------------------------------------------
    ws = wb["Rental Property"]
    lines = []
    for r in range(6, ws.max_row + 1):
        a = s(ws.cell(r, 1).value)
        if a:
            lines.append({"row": r, "desc": a, "control": s(ws.cell(r, 8).value) == "Control"})
    meta["labels"]["rental"] = lines

    # --- Queries / Review Notes numbered rows ------------------------------
    ws = wb["Queries"]
    meta["labels"]["query_rows"] = [r for r in range(8, ws.max_row + 1)
                                    if isinstance(ws.cell(r, 1).value, (int, float))]
    ws = wb["Review Notes"]
    meta["labels"]["review_rows"] = [r for r in range(8, ws.max_row + 1)
                                     if isinstance(ws.cell(r, 2).value, (int, float))]

    OUT.write_text(json.dumps(meta, indent=1, default=str), encoding="utf-8")
    print(f"wrote {OUT} ({OUT.stat().st_size:,} bytes): {len(meta['sheets'])} sheets, "
          f"{len(meta['external_names'])} external names, "
          f"{sum(len(v) for v in meta['baked_values'].values())} baked cells, "
          f"{len(rows)} summary rows, {len(blocks)} deduction blocks, {len(lines)} rental lines")


if __name__ == "__main__":
    main(*(sys.argv[1:2]))
